"""把 data/test_kb 下的文档灌进 Milvus。

用法：
    python -m scripts.ingest                 # 灌 data/test_kb
    python -m scripts.ingest <目录>          # 灌指定目录
    python -m scripts.ingest --query "问题"  # 只做一次检索验证
"""

from __future__ import annotations

import asyncio
import re
import sys
import time
from pathlib import Path

import pymupdf
from docx import Document

from agents.qa.cache import bump_cache_version
from agents.qa.retrieval.vector_store import get_vector_store
from agents.qa.splitter import split_document
from core.logger import get_logger

logger = get_logger(__name__)

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "test_kb"
SUFFIXES = {".md", ".txt", ".docx", ".pdf"}


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # 表格里的文字也要，否则会漏内容
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _read_pdf(path: Path) -> str:
    parts: list[str] = []
    with pymupdf.open(str(path)) as pdf:
        for page in pdf:
            text = page.get_text().strip()
            if not text:
                continue
            # PDF 里的换行多半是排版换行而非段落，中文直接接起来；
            # 空行（\n\n）才当作段落分隔。
            text = re.sub(r"(?<!\n)\n(?!\n)", "", text)
            parts.append(text)
    return "\n\n".join(parts)


_READERS = {
    ".md": _read_text,
    ".txt": _read_text,
    ".docx": _read_docx,
    ".pdf": _read_pdf,
}


def load_documents(directory: Path) -> list:
    docs = []
    for path in sorted(directory.iterdir()):
        suffix = path.suffix.lower()
        if suffix not in SUFFIXES:
            continue
        reader = _READERS[suffix]
        text = reader(path)
        if not text.strip():
            logger.warning("ingest.empty_document", file=path.name, suffix=suffix)
            continue
        source = path.stem
        chunks = split_document(text, source=source)
        docs.extend(chunks)
        logger.info(
            "ingest.file_loaded", file=path.name, suffix=suffix, chars=len(text), chunks=len(chunks)
        )
    return docs


def main() -> None:
    args = sys.argv[1:]

    if args and args[0] == "--query":
        query = args[1] if len(args) > 1 else "你们怎么收费"
        store = get_vector_store()
        results = store.hybrid_search_with_rerank(query)
        print(f"\n查询: {query}\n命中 {len(results)} 个父块\n")
        for i, doc in enumerate(results, 1):
            print(f"--- [{i}] source={doc.metadata.get('source')} ---")
            print(doc.page_content[:300].replace("\n", " "))
            print()
        return

    directory = Path(args[0]) if args and not args[0].startswith("-") else DATA_DIR
    if not directory.exists():
        raise SystemExit(f"目录不存在: {directory}")

    documents = load_documents(directory)
    if not documents:
        raise SystemExit(f"{directory} 下没有可导入的 {'/'.join(SUFFIXES)} 文件")

    print(f"切出 {len(documents)} 个子块，开始向量化并写入 Milvus…")
    store = get_vector_store()
    if "--reset" in args:
        print("重建集合（切分参数变了必须重建）…")
        store.reset()
    count = store.add_documents(documents)
    time.sleep(1.5)  # Milvus 写入后有短暂可见性延迟，立刻统计会得到 0
    print(f"写入 {count} 条，集合当前总条数: {store.count()}")

    # 知识库变了，旧答案缓存必须失效
    version = asyncio.run(bump_cache_version())
    print(f"缓存版本号已提升到 {version}，旧缓存失效")


if __name__ == "__main__":
    main()
